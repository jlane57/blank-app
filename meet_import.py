import csv
from copy import copy
import difflib
import re
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from spreadsheet_data import (
    find_year_sheet_layout,
    format_race_time,
    parse_meet_date,
    parse_race_time,
)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 500
SUPPORTED_EXTENSIONS = {".csv", ".tsv", ".xlsx", ".txt", ".pdf", ".png", ".jpg", ".jpeg"}
GROUPS = {"B": "Boys", "G": "Girls"}


def normalize_name(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def normalize_time(value):
    seconds = parse_race_time(value)
    if seconds is None:
        return ""
    return format_race_time(value)


def _clean(value):
    return str(value or "").strip()


def _rows_from_text(text, extension):
    if extension in {".csv", ".tsv"}:
        return list(csv.reader(StringIO(text), delimiter="\t" if extension == ".tsv" else ","))
    return [[line] for line in text.splitlines() if line.strip()]


def _rows_from_xlsx(data):
    workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        return [[cell.value for cell in row] for row in workbook.active.iter_rows()]
    finally:
        workbook.close()


def _rows_from_pdf(data):
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise ValueError("PDF support is not installed on this deployment.") from error
    reader = PdfReader(BytesIO(data))
    return [[line] for page in reader.pages for line in (page.extract_text() or "").splitlines()]


def _rows_from_image(data):
    try:
        import pytesseract
        from PIL import Image
    except ImportError as error:
        raise ValueError("Image OCR support is not installed on this deployment.") from error
    try:
        text = pytesseract.image_to_string(Image.open(BytesIO(data)))
    except Exception as error:
        raise ValueError("Image OCR could not run. Check the Tesseract installation.") from error
    return [[line] for line in text.splitlines()]


def extract_rows(filename, data):
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("The upload is larger than the 10 MB safety limit.")
    extension = Path(filename).suffix.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"{extension or 'This file type'} is not supported.")
    try:
        if extension == ".xlsx":
            rows = _rows_from_xlsx(data)
        elif extension == ".pdf":
            rows = _rows_from_pdf(data)
        elif extension in {".png", ".jpg", ".jpeg"}:
            rows = _rows_from_image(data)
        else:
            rows = _rows_from_text(data.decode("utf-8-sig"), extension)
    except (UnicodeDecodeError, ValueError, OSError) as error:
        raise ValueError(f"Could not read the uploaded file: {error}") from error
    rows = [[_clean(value) for value in row] for row in rows]
    if len(rows) > MAX_IMPORT_ROWS + 1:
        raise ValueError(f"The upload contains more than {MAX_IMPORT_ROWS} data rows.")
    return rows


def _header_indexes(rows):
    for row_index, row in enumerate(rows[:5]):
        indexes = {_clean(value).casefold(): index for index, value in enumerate(row) if _clean(value)}
        if any(key in indexes for key in ("name", "athlete", "runner")):
            return indexes, row_index
    return {}, 0


def _value(row, indexes, names):
    for name in names:
        index = indexes.get(name)
        if index is not None and index < len(row):
            return _clean(row[index])
    return ""


def extract_result_rows(rows):
    indexes, header_row = _header_indexes(rows)
    data_rows = rows[header_row + 1 :] if indexes else rows
    results = []
    for row_number, row in enumerate(data_rows, start=header_row + 2):
        if not any(row):
            continue
        if indexes:
            name = _value(row, indexes, ("name", "athlete", "runner"))
            time_value = _value(row, indexes, ("time", "result", "mark"))
        else:
            line = " ".join(row).strip()
            match = re.search(r"(?<!\d)(\d{1,3}:\d{2}(?:\.\d{1,3})?)\b", line)
            time_value = match.group(1) if match else ""
            name = line[: match.start()].strip(" ,-") if match else _clean(row[0] if row else "")
        if not time_value and name.casefold() in {"name", "athlete", "runner", "name time"}:
            continue
        normalized = normalize_time(time_value)
        if not name and not time_value:
            continue
        results.append(
            {
                "source_row": row_number,
                "source_name": name,
                "time": normalized or time_value,
                "valid_time": bool(normalized),
            }
        )
    if not results:
        raise ValueError("No athlete/time rows were found. Expected Name and Time columns.")
    return results


def match_to_roster(results, athletes, group):
    eligible = {
        normalize_name(athlete["name"]): athlete["name"]
        for athlete in athletes
        if athlete.get("team", "").strip().upper() == group
    }
    matched = []
    for result in results:
        normalized = normalize_name(result["source_name"])
        exact = eligible.get(normalized)
        candidates = difflib.get_close_matches(normalized, list(eligible), n=2, cutoff=0.72)
        if exact:
            athlete, confidence, status = exact, 1.0, "Exact"
        elif len(candidates) == 1:
            athlete = eligible[candidates[0]]
            confidence = difflib.SequenceMatcher(None, normalized, candidates[0]).ratio()
            status = "Review"
        else:
            athlete, confidence, status = "", 0.0, "Unmatched"
        matched.append(
            {
                **result,
                "athlete": athlete,
                "confidence": round(confidence, 2),
                "status": status if result["valid_time"] else "Invalid time",
            }
        )
    return matched


def _meet_columns(worksheet, layout):
    return [
        column
        for column in range(4, worksheet.max_column + 1)
        if _clean(worksheet.cell(layout["header_row"] + 1, column).value)
    ]


def _date_value(value):
    return parse_meet_date(value) or date.max


def _italic(cell, value):
    cell.value = str(value)
    font = copy(cell.font)
    font.italic = True
    cell.font = font


def apply_approved_results(
    workbook_data,
    year,
    group,
    meet_details,
    approved_rows,
    allow_overwrite=False,
):
    if group not in GROUPS:
        raise ValueError("Group must be B or G.")
    workbook = openpyxl.load_workbook(BytesIO(workbook_data))
    try:
        worksheet = workbook[str(year)]
        rows = [[cell.value for cell in row] for row in worksheet.iter_rows()]
        layout = find_year_sheet_layout(rows)
        if layout is None:
            raise ValueError(f"Year sheet {year} has no supported roster layout.")

        existing_name = meet_details.get("existing_meet", "")
        columns = _meet_columns(worksheet, layout)
        meet_column = None
        if existing_name:
            meet_column = next(
                (
                    column
                    for column in columns
                    if _clean(worksheet.cell(layout["header_row"] + 1, column).value) == existing_name
                ),
                None,
            )
            if meet_column is None:
                raise ValueError(f"Meet {existing_name!r} was not found.")
        else:
            meet_date = datetime.strptime(meet_details["date"], "%Y-%m-%d").date()
            meet_column = next(
                (
                    column
                    for column in columns
                    if _date_value(worksheet.cell(layout["date_row"] + 1, column).value) > meet_date
                ),
                worksheet.max_column + 1,
            )
            worksheet.insert_cols(meet_column)
            for row_index, value in {
                layout["header_row"] + 1: meet_details["name"],
                layout["meet_type_row"] + 1: meet_details["race_type"],
                layout["location_row"] + 1: meet_details["location"],
                layout["date_row"] + 1: meet_details["date"],
                layout["notes_row"] + 1: meet_details["notes"],
            }.items():
                _italic(worksheet.cell(row_index, meet_column), value)

        place_row = layout["boys_place_row"] if group == "B" else layout["girls_place_row"]
        if meet_details.get("team_place") and place_row is not None:
            _italic(worksheet.cell(place_row + 1, meet_column), meet_details["team_place"])

        roster_rows = {}
        for row in range(layout["first_athlete_row"] + 1, worksheet.max_row + 1):
            name = normalize_name(worksheet.cell(row, 1).value)
            team = _clean(worksheet.cell(row, 3).value).upper()
            if name and team == group:
                roster_rows[name] = row
        for item in approved_rows:
            athlete = normalize_name(item.get("athlete"))
            normalized_time = normalize_time(item.get("time"))
            if athlete not in roster_rows or not normalized_time:
                raise ValueError("Every approved row must have a valid athlete and time.")
            cell = worksheet.cell(roster_rows[athlete], meet_column)
            if cell.value not in (None, "") and not allow_overwrite:
                raise ValueError(f"{item['athlete']} already has a result; overwrite is disabled.")
            _italic(cell, normalized_time)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()
