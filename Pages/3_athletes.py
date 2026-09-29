from io import BytesIO
import re
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from zipfile import ZipFile
from xml.etree import ElementTree

import openpyxl
import pandas as pd
import requests
import streamlit as st

from site_status import stop_if_site_offline

stop_if_site_offline()

SPREADSHEET_ID = "1cO1crNvALaD4wS-T8Et5RDOSKLiHlWhbTT5vCCp-o3I"
MEET_FIRST_COLUMN = 4
MEET_LAST_COLUMN = 26
PLACEHOLDERS = {"", "n/a", "na", "---", "--", "-", "—", "–", "dnr", "dnf", "dq"}


def class_year(grade, roster_year):
	match = re.search(r"\d+", str(grade))
	if not match:
		return None

	grade_number = int(match.group())
	if not 9 <= grade_number <= 12:
		return None
	return roster_year + 13 - grade_number


def grade_standing(grade):
	match = re.search(r"\d+", str(grade))
	if not match:
		return ""

	return {
		9: "Freshman",
		10: "Sophomore",
		11: "Junior",
		12: "Senior",
	}.get(int(match.group()), "")


def _edit_distance(left, right):
	rows = len(left) + 1
	columns = len(right) + 1
	distance = [[0] * columns for _ in range(rows)]

	for i in range(rows):
		distance[i][0] = i
	for j in range(columns):
		distance[0][j] = j

	for i in range(1, rows):
		for j in range(1, columns):
			cost = 0 if left[i - 1] == right[j - 1] else 1
			distance[i][j] = min(
				distance[i - 1][j] + 1,
				distance[i][j - 1] + 1,
				distance[i - 1][j - 1] + cost,
			)
			if (
				i > 1
				and j > 1
				and left[i - 1] == right[j - 2]
				and left[i - 2] == right[j - 1]
			):
				distance[i][j] = min(
					distance[i][j], distance[i - 2][j - 2] + 1
				)

	return distance[-1][-1]


def _word_match_score(query_word, athlete_word):
	if query_word == athlete_word:
		return 1.0

	if athlete_word.startswith(query_word):
		return 0.82 + 0.16 * len(query_word) / len(athlete_word)
	if len(query_word) < 3 or len(athlete_word) < 3:
		return 0.0
	if query_word in athlete_word:
		return 0.88

	similarity = SequenceMatcher(None, query_word, athlete_word).ratio()
	allowed_edits = max(1, len(query_word) // 3)
	minimum_similarity = 0.66 if len(query_word) <= 4 else 0.55
	if (
		_edit_distance(query_word, athlete_word) <= allowed_edits
		and similarity >= minimum_similarity
	):
		return similarity
	return 0.0


def athlete_match_score(query, athlete, roster_years):
	query = str(query or "").strip().casefold()
	if not query:
		return 1.0

	query_years = [
		int(year) for year in re.findall(r"(?<!\d)\d{4}(?!\d)", query)
	]
	is_class_search = bool(re.search(r"\bclass\b", query))
	if query_years:
		if is_class_search:
			if athlete.get("class_year") not in query_years:
				return None
		elif not any(year in roster_years for year in query_years) and athlete.get("class_year") not in query_years:
			return None

	name_query = re.sub(
		r"\bclass\b|\bof\b|(?<!\d)\d{4}(?!\d)", " ", query
	)
	query_words = re.findall(r"[a-z0-9]+", name_query)
	if not query_words:
		return 1.0 if query_years else None
	if len(query_words) == 1 and len(query_words[0]) == 1 and not query_years:
		return None

	standings = athlete.get("search_standings") or [
		grade_standing(athlete.get("grade", ""))
	]
	searchable_text = " ".join(
		filter(None, [athlete.get("name", ""), *standings])
	).casefold()
	athlete_words = re.findall(r"[a-z0-9]+", searchable_text)
	name_words = re.findall(r"[a-z0-9]+", str(athlete.get("name", "")).casefold())
	scores = [
		max(
			(
				_word_match_score(word, candidate)
				for candidate in (name_words if len(word) <= 2 and len(query_words) > 1 else athlete_words)
			),
			default=0,
		)
		for word in query_words
	]
	if any(score == 0 for score in scores):
		return None

	match_score = sum(scores) / len(scores)
	if query_words == name_words:
		match_score += 0.2
	return match_score


def parse_time(value):
	text = str(value or "").strip()
	if not text:
		return None

	day_match = re.fullmatch(r"(\d+)\s+days?,?\s+(.+)", text, re.IGNORECASE)
	days = int(day_match.group(1)) if day_match else 0
	if day_match:
		text = day_match.group(2).strip()

	parts = text.split(":")
	if len(parts) in (2, 3):
		try:
			seconds = Decimal(parts[-1])
			if not seconds.is_finite() or not 0 <= seconds < 60:
				return None
			if len(parts) == 2:
				minutes = int(parts[0])
				if minutes < 0:
					return None
				return Decimal(days * 86400 + minutes * 60) + seconds
			hours, minutes = int(parts[0]), int(parts[1])
			if hours < 0 or not 0 <= minutes < 60:
				return None
			return Decimal(days * 86400 + hours * 3600 + minutes * 60) + seconds
		except (InvalidOperation, ValueError):
			return None

	try:
		day_fraction = Decimal(text)
		if day_fraction.is_finite() and 0 <= day_fraction <= 1:
			return day_fraction * Decimal(86400)
	except InvalidOperation:
		pass
	return None


def format_time(value):
	seconds = parse_time(value)
	if seconds is None:
		return str(value).strip()

	seconds = seconds.quantize(Decimal("0.001"))
	whole_seconds = int(seconds)
	minutes, remainder = divmod(whole_seconds, 60)
	result = f"{minutes:02d}:{remainder:02d}"
	fraction = seconds - Decimal(whole_seconds)
	if fraction:
		result += "." + f"{fraction:.3f}".split(".", 1)[1].rstrip("0")
	return result


@st.cache_data(ttl=3600)
def load_rosters():
	url = f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/export"
	response = requests.get(url, params={"format": "xlsx"}, timeout=30)
	response.raise_for_status()
	workbook_data = response.content

	with ZipFile(BytesIO(workbook_data)) as archive:
		root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
		sheet_names = [
			sheet.attrib["name"]
			for sheet in root.iter()
			if sheet.tag.rsplit("}", 1)[-1] == "sheet"
		]

	workbook = openpyxl.load_workbook(
		BytesIO(workbook_data), read_only=True, data_only=True
	)
	rosters = {}
	for name in sheet_names:
		if not re.fullmatch(r"\d{4}", name.strip()):
			continue

		year = int(name.strip())
		worksheet = workbook[name]
		meet_names = next(worksheet.iter_rows(
			min_row=2, max_row=2, min_col=MEET_FIRST_COLUMN,
			max_col=MEET_LAST_COLUMN, values_only=True,
		))
		meet_types = next(worksheet.iter_rows(
			min_row=3, max_row=3, min_col=MEET_FIRST_COLUMN,
			max_col=MEET_LAST_COLUMN, values_only=True,
		))
		athletes = []
		for row_number, values in enumerate(
			worksheet.iter_rows(min_row=5, min_col=1, max_col=MEET_LAST_COLUMN, values_only=True),
			start=5,
		):
			athlete_name = str(values[0] or "").strip()
			if not athlete_name:
				continue

			grade = str(values[1] or "").strip()
			results = []
			for index, result in enumerate(values[MEET_FIRST_COLUMN - 1:]):
				meet_name = str(meet_names[index] or "").strip()
				result_text = str(result or "").strip()
				if not meet_name or result_text.casefold() in PLACEHOLDERS:
					continue
				meet_type = str(meet_types[index] or "").strip()
				results.append({
					"year": year,
					"meet": meet_name,
					"type": meet_type,
					"raw_result": result,
					"is_5k": bool(re.search(
						r"(?<!\d)(?:5\s*k|5000(?:\s*(?:m|meter|metre))?)(?!\d)",
						f"{meet_name} {meet_type}",
						re.IGNORECASE,
					)),
				})
			athletes.append({
				"name": athlete_name,
				"grade": grade,
				"class_year": class_year(grade, year),
				"row": row_number,
				"results": results,
			})
		rosters[year] = athletes

	workbook.close()
	return rosters


st.title("Athletic Directory")

try:
	rosters = load_rosters()
except (requests.RequestException, OSError, KeyError, ValueError) as error:
	st.error(f"Could not load athlete rosters: {error}")
	st.stop()

for athletes in rosters.values():
	for athlete in athletes:
		for result in athlete["results"]:
			raw_result = result["raw_result"]
			result["result"] = format_time(raw_result)
			result["seconds"] = parse_time(raw_result)

years = sorted(rosters, reverse=True)
if not years:
	st.error("No yearly rosters were found in the spreadsheet.")
	st.stop()

latest_class_year = {}
for year in years:
	for item in rosters[year]:
		name = item["name"].casefold()
		if name not in latest_class_year:
			calculated_class_year = class_year(item["grade"], year)
			if calculated_class_year is not None:
				latest_class_year[name] = calculated_class_year

for year in years:
	for item in rosters[year]:
		item["class_year"] = latest_class_year.get(item["name"].casefold())

directory = {}
for year in years:
	for item in rosters[year]:
		name = item["name"].casefold()
		entry = directory.setdefault(name, {
			"year": year,
			"item": item,
			"roster_years": set(),
			"standings": set(),
		})
		entry["roster_years"].add(year)
		standing = grade_standing(item["grade"])
		if standing:
			entry["standings"].add(standing)

directory_by_year = {year: [] for year in years}
for entry in directory.values():
	directory_by_year[entry["year"]].append(entry)

selected = st.session_state.get("selected_athlete")
if selected:
	selected_name, selected_class_year = selected
	athlete = next(
		(
			item
			for year in years
			for item in rosters[year]
			if item["name"].casefold() == selected_name
			and item["class_year"] == selected_class_year
		),
		None,
	)
	if athlete is None:
		st.session_state.selected_athlete = None
		st.rerun()

	if st.button("← Back to all athletes"):
		st.session_state.selected_athlete = None
		st.rerun()

	athlete_results = [
		result
		for year in years
		for roster_athlete in rosters[year]
		if roster_athlete["name"].casefold() == selected_name
		and roster_athlete["class_year"] == selected_class_year
		for result in roster_athlete["results"]
	]
	athlete_results.sort(key=lambda result: result["year"], reverse=True)
	five_k_results = [
		result for result in athlete_results
		if result["is_5k"] and result["seconds"] is not None
	]
	if five_k_results:
		personal_record = min(five_k_results, key=lambda result: result["seconds"])
		personal_record_label = personal_record["result"]
	else:
		personal_record_label = "No 5K result"

	image_column, profile_column = st.columns([1, 3])
	with image_column:
		st.image("Images/raiders_logo.png", width=160, caption="Image placeholder")
	with profile_column:
		st.title(athlete["name"])
		class_label = f"Class of {athlete['class_year']}" if athlete["class_year"] else "Unknown"
		class_column, pr_column = st.columns(2)
		class_column.metric("Class", class_label)
		pr_column.metric("5K PR", personal_record_label)

	st.subheader("Meet times")
	if athlete_results:
		for year in sorted({result["year"] for result in athlete_results}, reverse=True):
			st.markdown(f"**{year}**")
			st.dataframe(
				pd.DataFrame([
					{
						"Meet": result["meet"],
						"Type": result["type"],
						"Time": result["result"],
					}
					for result in athlete_results
					if result["year"] == year
				]),
				height="content",
				hide_index=True,
				use_container_width=True,
			)
	else:
		st.info("No meet times were found for this athlete.")
else:
	search = st.text_input(
		"Search athletes",
		type="search",
		placeholder="Start typing an athlete name…",
		live=True,
	).strip().casefold()

	matches = []
	if search:
		for entry in directory.values():
			item = dict(entry["item"], search_standings=entry["standings"])
			match_score = athlete_match_score(search, item, entry["roster_years"])
			if match_score is not None:
				matches.append((match_score, entry["year"], entry["item"]))
		matches.sort(
			key=lambda match: (
				-match[0],
				-match[1],
				match[2]["name"].casefold(),
			)
		)

	def show_athlete_button(year, item, key):
		if st.button(item["name"], key=key, use_container_width=True):
			st.session_state.selected_athlete = (
				item["name"].casefold(), item["class_year"]
			)
			st.rerun()

	if search:
		if matches:
			st.caption("Matching athletes")
			for _, year, item in matches:
				show_athlete_button(
					year, item, f"search_{year}_{item['row']}"
				)
		else:
			st.info("No athletes match that search.")

	st.subheader("All athletes")
	for year in years:
		st.markdown(f"**{year}**")
		for item in rosters[year]:
			show_athlete_button(year, item, f"athlete_{year}_{item['row']}")