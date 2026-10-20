from base64 import b64encode
from pathlib import Path
from zipfile import BadZipFile

import requests
import streamlit as st
from openpyxl.utils.exceptions import InvalidFileException

from site_status import stop_if_site_offline
from spreadsheet_config import (
    get_spreadsheet_export_url,
    get_spreadsheet_id,
)
from spreadsheet_data import build_race_totals, parse_workbook_years
from gallery import get_gallery_entries

stop_if_site_offline()

try:
    featured_images = [
        image
        for image in get_gallery_entries()
        if image["category"].casefold() == "home"
        and image["featured"].casefold() == "true"
        and image["approved"].casefold() != "false"
    ][:6]
except (KeyError, OSError, ValueError, requests.RequestException):
    featured_images = []


@st.cache_data(ttl=3600)
def load_year_data(spreadsheet_id):
    response = requests.get(
        get_spreadsheet_export_url(spreadsheet_id),
        params={"format": "xlsx"},
        timeout=30,
    )
    response.raise_for_status()
    return parse_workbook_years(response.content)


hero_image = Path(__file__).resolve().parents[1] / "Images" / "group_huddle_home.jpg"
image_data = b64encode(hero_image.read_bytes()).decode("ascii")
st.markdown(
    f"""
    <style>
    [data-testid="stAppViewContainer"] > .main [data-testid="stMainBlockContainer"] {{
        width: 100% !important;
        max-width: none !important;
        padding-left: 0 !important;
        padding-right: 0 !important;
        padding-top: 0 !important;
    }}

    html,
    body,
    [data-testid="stAppViewContainer"] {{
        overflow-x: hidden;
    }}

    .stApp {{
        position: relative;
        isolation: isolate;
    }}

    .stApp::before {{
        content: "";
        position: fixed;
        inset: 0;
        z-index: 0;
        pointer-events: none;
        background-image:
            linear-gradient(
                0deg,
                rgba(255, 255, 255, 0.12),
                rgba(255, 255, 255, 0.12)
            ),
            url("data:image/jpeg;base64,{image_data}");
        background-size: cover;
        background-position: center;
        background-repeat: no-repeat;
    }}

    [data-testid="stAppViewContainer"] {{
        position: relative;
        z-index: 1;
    }}

    .home-hero {{
        display: grid;
        place-items: center;
        min-height: calc(100svh - 6rem);
        width: 100%;
        min-width: 0;
        box-sizing: border-box;
        margin: 0;
        padding: 1rem clamp(1rem, 3vw, 3rem) 2rem;
        position: relative;
    }}

    .home-hero::after {{
        content: "";
        position: absolute;
        bottom: 0;
        left: 50%;
        width: 100%;
        transform: translateX(-50%);
        border-bottom: 5px solid #F58220;
    }}

    .home-hero-content {{
        width: min(100%, 1800px);
        max-width: none;
        box-sizing: border-box;
        justify-self: center;
        margin-inline: auto;
        padding: clamp(2rem, 4vw, 3rem) clamp(2.5rem, 6vw, 5rem);
        border: 1px solid rgba(255, 255, 255, 0.7);
        border-top: 6px solid #F58220;
        border-radius: 0.5rem;
        background: rgba(255, 255, 255, 0.4);
        text-align: center;
        transform: translateY(-1rem);
    }}

    .home-hero-kicker {{
        margin: 0 0 0.8rem;
        color: #A54800;
        font-size: 0.78rem;
        font-weight: 800;
        letter-spacing: 0.16em;
        text-transform: uppercase;
    }}

    .home-hero h1 {{
        margin: 0 auto 1rem;
        max-width: 1000px;
        padding-bottom: 0;
        border-bottom: 0;
        color: #102A43;
        font-size: clamp(2.5rem, 6vw, 4.5rem);
        line-height: 1.12;
    }}

    .home-hero-divider {{
        width: 4.5rem;
        height: 4px;
        margin: 0 auto 1rem;
        border-radius: 999px;
        background: #F58220;
    }}

    .home-hero p {{
        margin: 0;
        color: #102A43;
        font-size: 1.15rem;
        font-weight: 600;
    }}

    .home-disclaimer {{
        margin: 0.5rem 0 1.25rem;
        color: rgba(16, 42, 67, 0.72);
        font-size: 0.72rem;
        text-align: center;
    }}

    .home-placeholders {{
        display: grid;
        grid-template-columns: repeat(2, minmax(0, 1fr));
        gap: 1.25rem;
        width: min(100%, 1040px);
        box-sizing: border-box;
        margin: 1.5rem auto 2rem;
        padding: 0 1rem;
    }}

    @media (min-width: 769px) {{
        body:has([data-testid="stSidebar"][aria-expanded="true"]) .stApp::before {{
            inset: 0 0 0 21rem;
        }}

        body:has([data-testid="stSidebar"][aria-expanded="true"]) .home-hero::after {{
            left: calc(50% + 10.5rem);
            width: calc(100% - 21rem);
        }}

        body:has([data-testid="stSidebar"][aria-expanded="true"]) .home-hero-content {{
            transform: translate(10.5rem, -1rem);
        }}

        body:has([data-testid="stSidebar"][aria-expanded="true"]) .home-disclaimer,
        body:has([data-testid="stSidebar"][aria-expanded="true"]) .home-placeholders {{
            position: relative;
            left: 10.5rem;
        }}
    }}

    .home-placeholder {{
        min-height: 12rem;
        border: 1px solid rgba(25, 57, 84, 0.28);
        border-top: 5px solid #F58220;
        border-radius: 0.5rem;
        background: rgba(255, 255, 255, 0.4);
        padding: 1.5rem;
        text-align: center;
    }}

    .home-stat-label {{
        color: #193954;
        font-size: 1.05rem;
        font-weight: 700;
    }}

    .home-stat-value {{
        margin: 0.4rem 0;
        color: #102A43;
        font-size: clamp(2rem, 4vw, 3rem);
        font-weight: 800;
        line-height: 1.1;
    }}

    .home-stat-note {{
        color: rgba(16, 42, 67, 0.72);
        font-size: 0.82rem;
    }}

    @media (max-width: 768px) {{
        [data-testid="stAppViewContainer"] > .main [data-testid="stMainBlockContainer"] {{
            padding-left: 1rem !important;
            padding-right: 1rem !important;
        }}

        .stApp {{
            background-color: #FFFFFF;
        }}

        .home-hero-content {{
            transform: translateY(-1rem);
        }}

        .home-disclaimer,
        .home-placeholders {{
            left: 0;
        }}

        .stApp::before {{
            display: none;
        }}

        .home-hero {{
            min-height: calc(100svh - 4rem);
            padding: 1rem 0.75rem 2rem;
            background-image:
                linear-gradient(
                    0deg,
                    rgba(255, 255, 255, 0.12),
                    rgba(255, 255, 255, 0.12)
                ),
                url("data:image/jpeg;base64,{image_data}");
            background-size: cover;
            background-position: center;
        }}

        .home-hero::after {{
            left: 50%;
            width: calc(100% - 1.5rem);
        }}

        .home-hero-content {{
            width: 100%;
        }}

        .home-placeholders {{
            grid-template-columns: 1fr;
        }}
    }}
    </style>
    <section class="home-hero">
        <div class="home-hero-content">
            <div class="home-hero-kicker">Boys' and Girls' Teams</div>
            <h1>Welcome to Harrison Cross Country!</h1>
            <div class="home-hero-divider"></div>
            <p>View results and stats for both teams</p>
        </div>
    </section>
    """,
    unsafe_allow_html=True,
)

if featured_images:
    st.markdown("### Featured Photos")
    photo_columns = st.columns(min(3, len(featured_images)))
    for index, image in enumerate(featured_images):
        with photo_columns[index % len(photo_columns)]:
            st.image(
                image["thumbnail_url"] or image["display_url"],
                caption=image["caption"] or None,
                use_container_width=True,
            )
            if image["photographer"]:
                st.caption(f"Photo: {image['photographer']}")

try:
    year_data = load_year_data(get_spreadsheet_id())
except requests.RequestException as error:
    st.error("Homepage information is temporarily unavailable.")
    totals = None
except (
    BadZipFile,
    OSError,
    ValueError,
    TypeError,
    KeyError,
    OverflowError,
    InvalidFileException,
) as error:
    st.error("Homepage information is temporarily unavailable.")
    totals = None
else:
    for issue in year_data.get("issues", []):
        st.warning(f"Skipped a year tab: {issue}")
    totals = build_race_totals(year_data["years"])

race_count = f"{totals['race_count']:,}" if totals is not None else "Unavailable"
hours_raced = f"{totals['hours_raced']:,.1f}" if totals is not None else "Unavailable"
since_year = (
    f"Since {totals['first_year']}"
    if totals is not None and totals["first_year"] is not None
    else "Start year unavailable"
)

st.markdown(
    f"""
    <div class="home-placeholders" aria-label="Reserved homepage content areas">
        <section class="home-placeholder" aria-label="Total races run">
            <div class="home-stat-label">Total Races Run</div>
            <div class="home-stat-value">{race_count}</div>
            <div class="home-stat-note">{since_year}</div>
        </section>
        <section class="home-placeholder" aria-label="Total hours raced">
            <div class="home-stat-label">Total Hours Raced</div>
            <div class="home-stat-value">{hours_raced}</div>
            <div class="home-stat-note">{since_year}</div>
        </section>
    </div>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    "<div class='home-disclaimer'>Photo credit River Bender.</div>",
    unsafe_allow_html=True,
)

st.markdown(
    "<div class='home-disclaimer'>This site is not endorsed by Harrison High School and is intended for internal use only.</div>",
    unsafe_allow_html=True,
)